# ai_reports/utils.py
import os
import PyPDF2
import docx
from io import BytesIO

def extract_text_from_file(file_path):
    """Extract text from a file (PDF or DOCX)"""
    if not os.path.exists(file_path):
        return ""
        
    if file_path.lower().endswith('.pdf'):
        return _extract_text_from_pdf(file_path)
    elif file_path.lower().endswith(('.doc', '.docx')):
        return _extract_text_from_docx(file_path)
    else:
        return ""

def extract_text_from_file_object(file_obj):
    """Extract text from a file object (PDF or DOCX)"""
    if not file_obj:
        return ""
        
    file_content = file_obj.read()
    file_obj.seek(0)  # Reset file pointer
    
    # Determine file type from file object
    if hasattr(file_obj, 'name'):
        file_name = file_obj.name.lower()
        if file_name.endswith('.pdf'):
            return _extract_text_from_pdf_content(file_content)
        elif file_name.endswith(('.doc', '.docx')):
            return _extract_text_from_docx_content(file_content)
    
    return ""

def _extract_text_from_pdf(file_path):
    """Extract text from a PDF file"""
    with open(file_path, 'rb') as file:
        pdf_reader = PyPDF2.PdfReader(file)
        text = ""
        for page in pdf_reader.pages:
            text += page.extract_text()
        return text

def _extract_text_from_pdf_content(file_content):
    """Extract text from PDF content"""
    pdf_reader = PyPDF2.PdfReader(BytesIO(file_content))
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

def _extract_text_from_docx(file_path):
    """Extract text from a DOCX file"""
    doc = docx.Document(file_path)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def _extract_text_from_docx_content(file_content):
    """Extract text from DOCX content"""
    doc = docx.Document(BytesIO(file_content))
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text